"""Inspeção visual de MT no mapa do alimentador: lê a base de inspeção (OS INSP-RD, aba "Dados") e grava
alimentadores/<AL>/inspecao.js com as inspeções EXECUTADAS e os vãos que cada uma cobriu.

Onde a inspeção passou: a rede do ELEMENTO subindo até a ABRANGENCIA. Conferido na base de 13/10/2026: nas 1.170 OS
com os dois ativos no mapa e em postes diferentes, a ABRANGENCIA está sempre a montante do ELEMENTO (e a
MONT_ABRANGENCIA, a montante da ABRANGENCIA); as OS seguidas formam a corrida pela rede (a ABRANGENCIA de uma é o
ELEMENTO da outra).
  - ELEMENTO e ABRANGENCIA no mapa do alimentador da OS: os vãos do caminho entre os dois.
  - Os dois no mesmo poste, ou só um deles no mapa: só o ponto (a OS aparece na lista e, ao tocar, marca o ponto).
    O outro ativo não está no cadastro do site: em geral os "90..." (bloco de chave faca, sem ponto no GIS nem no KML).
  - Nenhum dos dois: sem posição (fica na lista).
Ativo no mapa: chave do alimentador (dados/alimentadores/<AL>.json) ou, pelo índice de ativos do site (dados/ativos:
transformadores e chaves com coordenada), o poste do mapa mais perto, até LIM_M.
Data da inspeção: FIM_EXEC; se o FIM vier mais de 60 dias depois do INICIO_EXEC (fechamento em lote), INICIO_EXEC
(a mesma regra da limpeza de faixa).
Todo alimentador com página ganha o inspecao.js (sem OS quando não tem inspeção): a página não pede arquivo que não existe.

Uso:
  python scripts/inspecao.py "BASE INSPEÇÃO.xlsx" <pasta do site>
Depois, rodar gera_trechos.py --so-interface (a página do alimentador carrega o inspecao.js).
"""
import datetime as dt
import json
import math
import re
import sys
from collections import Counter, defaultdict, deque
from pathlib import Path

import openpyxl

sys.path.insert(0, str(Path(__file__).resolve().parent))
from gera_trechos import chave_ativo, fatia  # noqa: E402  (mesma chave e mesma fatia do índice de ativos)

LIM_M = 50          # ativo do índice a mais que isso do poste mais perto do mapa não entra (mesmo limite dos trafos)
LOTE_DIAS = 60      # FIM_EXEC mais de 60 dias depois do INICIO_EXEC: vale o INICIO_EXEC
SEM_DESC = ('OS despachada automaticamente',)   # descrição automática do sistema: não vai para a página


def le_js(p):
    s = p.read_text(encoding='utf-8')
    return json.loads(s[s.index('=') + 1:].strip().rstrip(';'))


def le_base(xlsx):
    """Linhas da aba Dados (ou da 1ª aba) como dicionários."""
    wb = openpyxl.load_workbook(xlsx, read_only=True, data_only=True)
    ws = wb['Dados'] if 'Dados' in wb.sheetnames else wb.worksheets[0]
    linhas = ws.iter_rows(values_only=True)
    cab = [str(c).strip() if c is not None else '' for c in next(linhas)]
    return [dict(zip(cab, r)) for r in linhas if any(v not in (None, '') for v in r)]


def texto(v):
    return '' if v is None else re.sub(r'\s+', ' ', str(v)).strip()


def data_insp(r):
    ini, fim = r.get('INICIO_EXEC'), r.get('FIM_EXEC')
    if not isinstance(fim, dt.datetime):
        return None
    if isinstance(ini, dt.datetime) and (fim - ini).days > LOTE_DIAS:
        return ini.date().isoformat()
    return fim.date().isoformat()


class Indice:
    """dados/ativos/<fatia>.json do site, lido por pedaço conforme a busca."""
    def __init__(self, site):
        self.pasta, self.cache = site / 'dados' / 'ativos', {}

    def busca(self, cod):
        k = chave_ativo(cod)
        f = fatia(k)
        if f not in self.cache:
            p = self.pasta / f'{f}.json'
            self.cache[f] = json.loads(p.read_text(encoding='utf-8')) if p.exists() else {}
        return self.cache[f].get(k, [])


class Rede:
    """Mapa do alimentador (dados.js): árvore a partir do DJ, para o caminho entre dois postes."""
    def __init__(self, site, al):
        self.D = le_js(site / 'alimentadores' / al / 'dados.js')
        js = json.loads((site / 'dados' / 'alimentadores' / f'{al}.json').read_text(encoding='utf-8'))
        self.chave = {chave_ativo(c['nome']): c['poste'] for c in js['chaves'] if c.get('poste')}
        viz, self.coord = defaultdict(list), {}
        for i, v in enumerate(self.D['vaos']):
            viz[v[6]].append((v[7], i))
            viz[v[7]].append((v[6], i))
            self.coord.setdefault(v[6], (v[0], v[1]))
            self.coord.setdefault(v[7], (v[2], v[3]))
        raiz = self.D['trechos'][0]['inicio_poste']
        self.pai, self.nivel = {raiz: (None, None)}, {raiz: 0}
        fila = deque([raiz])
        while fila:
            p = fila.popleft()
            for q, i in viz[p]:
                if q not in self.pai:
                    self.pai[q], self.nivel[q] = (p, i), self.nivel[p] + 1
                    fila.append(q)

    def poste_perto(self, lat, lon):
        k = math.cos(math.radians(lat))
        melhor, dmin = None, float('inf')
        for p, (la, lo) in self.coord.items():
            d = ((lo - lon) * k) ** 2 + (la - lat) ** 2
            if d < dmin:
                melhor, dmin = p, d
        return melhor, math.sqrt(dmin) * 111320

    def caminho(self, a, b):
        """Vãos entre os postes a e b pela árvore (None se algum não liga ao DJ)."""
        if a not in self.pai or b not in self.pai:
            return None
        vs = []
        while a != b:
            if self.nivel[a] >= self.nivel[b]:
                a, i = self.pai[a]
            else:
                b, i = self.pai[b]
            vs.append(i)
        return sorted(vs)


def main(xlsx, site):
    site = Path(site)
    base = le_base(xlsx)
    exec_ = [r for r in base if texto(r.get('SITUACAO')).upper() == 'EXECUTADA']
    datas = [d for d in map(data_insp, exec_) if d]
    pags = sorted(p.parent.name for p in (site / 'alimentadores').glob('*/dados.js')
                  if (p.parent / 'index.html').exists() and (site / 'dados' / 'alimentadores' / f'{p.parent.name}.json').exists())
    idx, redes = Indice(site), {}
    por_al = {a: [] for a in pags}
    cont, sem_pagina = Counter(), Counter()

    def ponto(rede, al, cod):
        """código -> poste do mapa do alimentador (ou None)."""
        k = chave_ativo(cod)
        if not k:
            return None
        if k in rede.chave:
            return rede.chave[k]
        for a, _tipo, _t, _papel, lat, lon in idx.busca(cod):
            if a == al and lat is not None:
                p, d = rede.poste_perto(lat, lon)
                if d <= LIM_M:
                    return p
        return None

    for r in exec_:
        al = texto(r.get('ALIMENTADOR')).upper()
        if al not in por_al:
            sem_pagina[al] += 1
            continue
        rede = redes.get(al) or redes.setdefault(al, Rede(site, al))
        el, ab, mo = (texto(r.get(c)) for c in ('ELEMENTO', 'ABRANGENCIA', 'MONT_ABRANGENCIA'))
        pe, pa = ponto(rede, al, el), ponto(rede, al, ab)
        vs, pt = [], None
        if pe and pa and pe != pa and (c := rede.caminho(pe, pa)) is not None:
            vs, metodo = c, 'entre o ELEMENTO e a ABRANGÊNCIA'
        elif pe and pa:
            pt, metodo = pe, 'ELEMENTO e ABRANGÊNCIA no mesmo ponto'
        elif pe:
            pt, metodo = pe, f'só o ponto do ELEMENTO: a ABRANGÊNCIA {ab or "(vazia)"} não está no mapa deste alimentador'
        elif pa:
            pt, metodo = pa, f'só o ponto da ABRANGÊNCIA: o ELEMENTO {el or "(vazio)"} não está no mapa deste alimentador'
        else:
            metodo = 'sem posição: ELEMENTO e ABRANGÊNCIA não estão no mapa deste alimentador'
        cont['no mapa (vãos)' if vs else 'só o ponto' if pt else 'sem posição'] += 1
        num = texto(r.get('NUMERO_OS'))
        desc = texto(r.get('DESCRICAO_OS'))
        reg = {
            'os': num.split()[-1] if num else '',
            'os_completa': num,
            'eq': texto(r.get('CODEQUIPE')),
            'data': data_insp(r),
            'crit': texto(r.get('CRITICIDADE')).lower(),
            'elem': el, 'abr': ab, 'mont': mo,
            'desc': '' if desc.startswith(SEM_DESC) else desc.replace('�', '?')[:300],
            'metodo': metodo,
            'm': round(sum(rede.D['vaos'][i][8] for i in vs), 1),
            'pt': list(rede.coord[pt]) if pt else None,
        }
        por_al[al].append((reg, vs))

    hoje = dt.date.today().isoformat()
    km = 0.0
    for al, lst in por_al.items():
        lst.sort(key=lambda ov: (ov[0]['data'] or '', ov[0]['os_completa']))
        out = {'fonte': Path(xlsx).name, 'gerado': hoje, 'base_ate': max(datas) if datas else None,
               'os': [o for o, _ in lst], 'vaos': {}}
        if lst:
            rede = redes[al]
            for k, (_, vs) in enumerate(lst):
                for i in vs:
                    out['vaos'].setdefault(rede.D['vaos'][i][5], []).append(k)
            km += sum(rede.D['vaos'][i][8] for i in {i for _, vs in lst for i in vs}) / 1000
        (site / 'alimentadores' / al / 'inspecao.js').write_text(
            'window.INSPECAO = ' + json.dumps(out, ensure_ascii=False, separators=(',', ':')) + ';\n', encoding='utf-8')
    return dict(linhas=len(base), executadas=len(exec_), em_paginas=sum(cont.values()), situacao=dict(cont),
                km_inspecionados=round(km, 1), alimentadores_com_inspecao=sum(bool(v) for v in por_al.values()),
                paginas=len(pags), sem_pagina=dict(sem_pagina.most_common()), base_ate=max(datas) if datas else None)


if __name__ == '__main__':
    if len(sys.argv) != 3:
        sys.exit(__doc__)
    r = main(sys.argv[1], sys.argv[2])
    print(f"{r['linhas']} OS na base · {r['executadas']} executadas (até {r['base_ate']}) · "
          f"{r['em_paginas']} em alimentadores do site")
    for k, n in r['situacao'].items():
        print(f'  {k}: {n}')
    print(f"  {r['km_inspecionados']} km de rede inspecionados em {r['alimentadores_com_inspecao']} alimentadores "
          f"(inspecao.js gravado nos {r['paginas']} com página)")
    if r['sem_pagina']:
        print('  alimentador sem página no site (fora do mapa): ' +
              ', '.join(f'{a} ({n})' for a, n in r['sem_pagina'].items()))
